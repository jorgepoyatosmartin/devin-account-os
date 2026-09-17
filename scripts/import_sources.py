"""Import MAPFRE Territory Plan (map mapfre.pptx, March 2026) and ACCIONA PG target list (Acciona.xlsx).
Source files are treated as source of truth; public research is kept as secondary, never overwrites."""
import json, pathlib, re, unicodedata
import openpyxl
from pptx import Presentation

ROOT = pathlib.Path(__file__).resolve().parents[1] / 'src' / 'data'
PPTX = '/home/ubuntu/attachments/5e967fdf-dbe1-4a72-a61b-3f2e6a74b2f1/map_mapfre.pptx'
XLSX = '/home/ubuntu/attachments/a2c83032-2572-4752-943c-d9ed62adee99/Acciona.xlsx'
TODAY = '2026-09-17'
UNK = 'UNKNOWN'
TP_SRC = 'MAPFRE Territory Plan FY27 — Power Chart slide (map mapfre.pptx, March 2026)'
XL_SRC = 'Acciona.xlsx — sheet "PG", PG TARGET table'

stk = json.load(open(ROOT / 'stakeholders.json'))
S = {s['id']: s for s in stk}


def slug(name):
    n = unicodedata.normalize('NFKD', name).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z]+', '-', n).strip('-')


# ------------------------------------------------------------------ MAPFRE
prs = Presentation(PPTX)
slide = prs.slides[1]
boxes = []
for sh in slide.shapes:
    if sh.shape_type == 1 and sh.has_text_frame and sh.text_frame.text.strip():
        lines = [l.strip() for l in sh.text_frame.text.split('\n') if l.strip()]
        boxes.append(dict(name=lines[0], title=lines[1] if len(lines) > 1 else UNK,
                          role=(lines[2] if len(lines) > 2 else '-'), x=sh.left + sh.width // 2, y=sh.top, bottom=sh.top + sh.height))

# reporting lines: derived from connector geometry on the slide (verified manually)
REPORTS = {
    'Countries': 'Vanessa Escrivá', 'Santiago Wiznez': 'Vanessa Escrivá', 'José Luís Bernal': 'Vanessa Escrivá',
    'José Antonio Ledesma': 'Vanessa Escrivá', 'Maribel Solanas': 'Vanessa Escrivá', 'Leire Jiménez': 'Vanessa Escrivá',
    'Íñigo Lacave': 'Santiago Wiznez', 'Javier Maraña': 'Santiago Wiznez', 'Jesús López': 'Santiago Wiznez',
    'Enrique Turillo': 'José Antonio Ledesma', 'DIEGO JOSÉ BODAS': 'Maribel Solanas',
    'Mat Javanovic': 'Íñigo Lacave', 'Gabriel Andújar': 'Javier Maraña', 'Mario Encinar': 'DIEGO JOSÉ BODAS',
    'Andrés Hevia': 'Mat Javanovic', 'Gonzalo Cabanillas': 'Gabriel Andújar',
}
TITLES = {'TECHNOLOGY & ARCHITECTURE DIRECTOR': 'Technology & Architecture Director', 'GLOBAL CLOUD DIRECTOR': 'Global Cloud Director', 'CTTO': 'CTTO', 'CISO': 'CISO', 'CDO': 'CDO',
          'HEAD OF AGENTIC ARCHITECTURE & TECHNOLOGICAL INNOVATION': 'Head of Agentic Architecture & Technological Innovation', 'IT ARCHITECTURE': 'IT Architecture',
          'AGENT ARCHITECTURE MANAGER': 'Agent Architecture Manager', 'SENIOR SOFTWARE ENGINEER': 'Senior Software Engineer', 'DATA AND IA CIBER DIRECTOR': 'Data and IA Ciber Director',
          'AI CENTER DIRECTOR': 'AI Center Director', 'AiI LEAD EXPERT': 'AI Lead Expert', 'HTECH INNOVATION MANAGER': 'HTech Innovation Manager', 'CIO S': 'Country CIOs'}
ROLE = {'EB': 'Economic Buyer', 'CHAMPION': 'Champion', 'T. CHAMPION': 'Technical Champion', 'COACH': 'Coach', '-': 'Unknown'}
LEVEL_BY_ROW = [(1000000, 'Executive Committee'), (1800000, 'Senior Leadership'), (2600000, 'Director'), (3400000, 'Manager'), (9e9, 'Individual Contributor')]
EXISTING = {'Vanessa Escrivá': 'mapfre-escriva', 'José Luís Bernal': 'mapfre-bernal', 'Santiago Wiznez': 'mapfre-wiznez', 'José Antonio Ledesma': 'mapfre-ledesma',
            'Maribel Solanas': 'mapfre-solanas', 'Leire Jiménez': 'mapfre-jimenez', 'Íñigo Lacave': 'mapfre-lacave', 'Mat Javanovic': 'mapfre-javanovic',
            'Javier Maraña': 'mapfre-marana', 'Gabriel Andújar': 'mapfre-andujar', 'DIEGO JOSÉ BODAS': 'mapfre-bodas'}
ids = {n: EXISTING.get(n, 'mapfre-' + slug(n)) for n in REPORTS} | {'Vanessa Escrivá': 'mapfre-escriva'}
ids['Countries'] = 'mapfre-countries-cios'

def pretty(n):
    return 'Diego José Bodas' if n == 'DIEGO JOSÉ BODAS' else n

for b in boxes:
    sid = ids[b['name']]
    role = ROLE[b['role']]
    level = next(l for y, l in LEVEL_BY_ROW if b['y'] < y)
    if b['name'] == 'Vanessa Escrivá':
        level = 'Executive Committee'
    title_tp = TITLES.get(b['title'], b['title'])
    ev = dict(claim=f"{pretty(b['name'])} — {title_tp}" + (f" — {role}" if role != 'Unknown' else ''), source=TP_SRC, date='2026-03', category='FACT', confidence='High')
    rel = {'Champion': 'Champion', 'Technical Champion': 'Champion', 'Coach': 'Engaged'}.get(role, 'Identified')
    s = S.get(sid)
    if s is None:
        s = dict(id=sid, accountId='mapfre', name=pretty(b['name']) if b['name'] != 'Countries' else 'Country CIOs (group)', title=title_tp,
                 functionArea=title_tp, responsibilities=UNK, strategicPriorities=UNK, technologyPriorities=UNK, relevantInitiatives=[],
                 publicStatements=[], recentActivity=UNK, potentialPain=UNK, cognitionRelevance=UNK, relationshipStatus=rel, buyingRole=role,
                 sources=[], useCaseIds=[], influence='Unknown', championPotential='Unknown')
        stk.append(s); S[sid] = s
    else:
        # territory plan title wins; keep public-research title in a note
        pub = s['title'].split(' (Territory Plan)')[0] if '(Territory Plan)' in s['title'] else s['title']
        if 'PUBLIC RESEARCH' in s['title']:
            pub = s['title'].split('· ')[-1].replace(' (PUBLIC RESEARCH)', '')
        s['title'] = title_tp if pub.lower().startswith(title_tp.lower()) else f'{title_tp} · public: {pub}'
    s['sources'] = [e for e in s.get('sources', []) if 'user brief' not in e['source']] + [ev]
    keep_hyp = role == 'Unknown' and s.get('powerRole') not in (None, 'Unknown')
    if keep_hyp:
        role = s['powerRole']
    s.update(powerRole=role, roleIsHypothesis=keep_hyp, buyingRole=role, level=level, dataOrigin='TERRITORY PLAN', lastUpdated=TODAY,
             relationshipStatus=s['relationshipStatus'] if s.get('relationshipStatus') in ('Champion', 'Engaged') else rel,
             championPotential='High' if role in ('Champion', 'Technical Champion', 'Coach') else s.get('championPotential', 'Unknown'),
             recommendedNextAction='Confirm current title and agenda via the Coach (Íñigo Lacave / Mario Encinar) before outreach' if s.get('recommendedNextAction', '').startswith('Confirm') else s.get('recommendedNextAction', UNK))
    if b['name'] in REPORTS:
        s['reportsTo'] = ids[REPORTS[b['name']]]
    s['territoryPlan'] = dict(role=b['role'], title=b['title'], reportsTo=pretty(REPORTS.get(b['name'], '')) or None, source=TP_SRC)

# Malu Delicado (public research, CIO Iberia) fits the plan's "Country CIOs" box → interpretation, not fact
S['mapfre-delicado']['reportsTo'] = 'mapfre-countries-cios'
S['mapfre-delicado']['sources'].append(dict(claim='Grouped under "Countries CIOs" box of the Territory Plan power chart (interpretation — not named in the deck)', source=TP_SRC, date='2026-03', category='SOURCE-BASED INTERPRETATION', confidence='Medium'))
S['mapfre-escriva'].pop('reportsTo', None)

# ------------------------------------------------------------------ ACCIONA
wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb.worksheets[0]
rows = [r for r in ws.iter_rows(min_row=3, values_only=True) if r[3] and isinstance(r[2], (int, float))]
XL_LEVEL = {'C Level': 'Executive Committee', '2nd': 'Director'}
XL_MATCH = {'Raúl Rivero': 'ana-rivero', 'José Manuel Carballo': 'ana-carballo'}
def clean(v):
    return (str(v).replace('\u200b', '').replace('·', '').strip()) if v is not None else ''
for r in rows:
    person, title, level, bu, play, action, why = (clean(x) for x in r[3:10])
    excel = dict(person=person, title=title, level=level, businessUnit=bu, salesPlay=play or UNK, action=action or UNK, whyHighTarget=why or UNK)
    sid = XL_MATCH.get(person, 'ana-' + slug(person))
    ev = dict(claim=f'{person} — {title} — {level} — {bu} (PG target #{int(r[2])})', source=XL_SRC, date='2026', category='FACT', confidence='High')
    s = S.get(sid)
    if s is None:
        s = dict(id=sid, accountId='acciona', name=person, title=title, functionArea=bu, responsibilities=UNK, strategicPriorities=UNK,
                 technologyPriorities=UNK, relevantInitiatives=[], publicStatements=[], recentActivity=UNK, potentialPain=UNK, cognitionRelevance=UNK,
                 relationshipStatus='Identified', buyingRole='Unknown', sources=[], powerRole='Unknown', roleIsHypothesis=False, useCaseIds=[],
                 influence='Unknown', championPotential='Unknown', recommendedNextAction='Fill Sales Play / Action / Why High Target in the PG sheet, then plan first touch')
        stk.append(s); S[sid] = s
    else:
        if s['title'] != title:
            s['title'] = f"{title} · public: {s['title']}"
    s['sources'].append(ev)
    s.update(excel=excel, dataOrigin='ORIGINAL EXCEL DATA', level=XL_LEVEL.get(level, 'Unknown'), businessUnit=bu, lastUpdated=TODAY)

json.dump(stk, open(ROOT / 'stakeholders.json', 'w'), ensure_ascii=False, indent=2); open(ROOT / 'stakeholders.json', 'a').write('\n')
print(len(stk), 'stakeholders;', sum(s['accountId'] == 'mapfre' for s in stk), 'mapfre;', sum(s['accountId'] == 'acciona' for s in stk), 'acciona')
